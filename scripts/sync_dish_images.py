#!/usr/bin/env python3
"""
Sync Dish Images Script
-----------------------
Scans the `dish_photos/` folder for authentic dish photos, matches each photo
with dishes in `data/dishes.csv` (fuzzy & exact matching), copies them to
`static/images/dishes/`, updates `data/dishes.csv`, generates `data/dishes.xlsx` (Excel),
and updates `kitchen.db`!
"""

import os
import re
import csv
import shutil
from pathlib import Path

# Paths
BASE_DIR = Path(__file__).resolve().parent.parent
PHOTOS_DIR = BASE_DIR / "dish_photos"
STATIC_DISHES_DIR = BASE_DIR / "static" / "images" / "dishes"
CSV_PATH = BASE_DIR / "data" / "dishes.csv"
EXCEL_PATH = BASE_DIR / "data" / "dishes.xlsx"
DB_PATH = BASE_DIR / "kitchen.db"

SUPPORTED_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".avif", ".gif"}


def clean_name(s: str) -> str:
    """Normalize string for fuzzy comparison."""
    s = s.lower()
    s = re.sub(r'[\-_]+', ' ', s)
    s = re.sub(r'[^a-z0-9 ]', '', s)
    return " ".join(s.split())


def sync_images():
    print("=" * 60)
    print("🍳 Chef Chat Dish Photos & Excel Synchronizer")
    print("=" * 60)

    PHOTOS_DIR.mkdir(parents=True, exist_ok=True)
    STATIC_DISHES_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Load CSV
    if not CSV_PATH.exists():
        print(f"❌ Error: {CSV_PATH} not found.")
        return

    with open(CSV_PATH, mode="r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
        fieldnames = reader[0].keys() if reader else []

    if not reader:
        print("❌ dishes.csv is empty.")
        return

    print(f"📖 Loaded {len(reader)} dishes from CSV.")

    # 2. Scan dish_photos/ folder
    photo_files = [
        f for f in PHOTOS_DIR.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTS
    ]
    print(f"📸 Found {len(photo_files)} photos in {PHOTOS_DIR.name}/")

    matched_count = 0
    updated_dishes = []

    # Map dishes by normalized name
    dish_map = {clean_name(row["name"]): row for row in reader}

    for photo in photo_files:
        photo_stem = clean_name(photo.stem)
        matched_dish = None

        # Exact match
        if photo_stem in dish_map:
            matched_dish = dish_map[photo_stem]
        else:
            # Fuzzy sub-match (e.g. "chicken karahi" matches "Chicken Karahi Desi Style")
            for norm_name, row in dish_map.items():
                if photo_stem in norm_name or norm_name in photo_stem:
                    matched_dish = row
                    break

        if matched_dish:
            ext = photo.suffix.lower()
            clean_filename = re.sub(r'[^a-zA-Z0-9_\-]', '_', matched_dish['name']).lower() + ext
            dest_path = STATIC_DISHES_DIR / clean_filename
            
            # Copy to static/images/dishes/
            shutil.copy2(photo, dest_path)
            rel_image_path = f"images/dishes/{clean_filename}"
            matched_dish["image"] = rel_image_path
            matched_count += 1
            print(f"  ✓ Matched '{photo.name}' -> Dish: '{matched_dish['name']}' -> {rel_image_path}")
        else:
            print(f"  ⚠️ No matching dish found for photo: '{photo.name}' (Tip: Rename file closer to dish name)")

    # 3. Save updated CSV
    with open(CSV_PATH, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(reader)
    print(f"💾 Updated {CSV_PATH}")

    # 4. Generate Excel Sheet (.xlsx)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Chef Chat Dishes"

        headers = list(fieldnames)
        ws.append(headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="CC4C39", end_color="CC4C39", fill_type="solid") # Brand terracotta
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, row_data in enumerate(reader, 2):
            for col_idx, key in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                if key in ("type", "course", "cuisine"):
                    cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        wb.save(EXCEL_PATH)
        print(f"📊 Created formatted Excel Spreadsheet: {EXCEL_PATH}")
    except Exception as e:
        print(f"⚠️ Notice: Could not generate Excel: {e}")

    # 5. Update SQLite Database if dishes exist there
    if DB_PATH.exists():
        try:
            import sqlite3
            conn = sqlite3.connect(DB_PATH)
            cur = conn.cursor()
            for r in reader:
                cur.execute(
                    "UPDATE custom_dishes SET image = ? WHERE LOWER(name) = LOWER(?)",
                    (r["image"], r["name"])
                )
            conn.commit()
            conn.close()
            print("🗄️ Synchronized SQLite database custom dishes.")
        except Exception as e:
            print(f"Database sync note: {e}")

    print("=" * 60)
    print(f"✅ Sync complete! {matched_count} images linked successfully.")
    print(f"📂 Drop more images anytime into: {PHOTOS_DIR}")
    print("=" * 60)


if __name__ == "__main__":
    sync_images()
